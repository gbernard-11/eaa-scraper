#!/usr/bin/env python3
"""Export EAA Young Eagles final-report data for a date range to Excel."""

from __future__ import annotations

import json
import re
import sys
import time
from collections import defaultdict
from datetime import date, datetime
from getpass import getpass
from pathlib import Path
from typing import Any

import requests
from dateutil import parser as date_parser
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

API_BASE = "https://events.eaachapters.org/web/api/"
SITE_ORIGIN = "https://www.eaachapters.org"
LOGIN_URL = f"{SITE_ORIGIN}/login"
RECAPTCHA_SITE_KEY = "6LeyIN0pAAAAAHKShVnVIp6sy9Yk7MgEzhvdm_zK"
REQUEST_PAUSE_SEC = 0.35
EVENT_PAGE_SIZE = 10
REPORT_PAGE_SIZE = 50
REPORTS_DIR = Path(__file__).resolve().parent / "Reports"

STATUS_COMPLETED = 5
STATUS_CANCELLED = 6
STATUS_NAMES = {
    1: "INCOMPLETE",
    2: "CREATED",
    3: "OPEN",
    4: "CLOSED",
    5: "COMPLETED",
    6: "CANCELLED",
    7: "WAITLIST",
    8: "SYNCHRONIZING",
    9: "WALKUP",
}

PARENT_FILTERS = (
    (1, "Parents Registered", "Registered parents"),
    (2, "Parents Flown", "Parents of youth flown"),
    (3, "Parents Waitlisted", "Parents of youth waitlisted"),
    (4, "Parents Not Flown", "Parents of youth that didn't fly"),
)

HQ_LABELS = {
    "totalYoungs": "Registered Youth",
    "totalGirls": "Registered Females",
    "totalBoys": "Registered Males",
    "totalNS": "Registered Undisclosed Gender",
    "totalYoungsFlown": "Youth Flown",
    "totalGirlsFlown": "Females Flown",
    "totalBoysFlown": "Males Flown",
    "totalNSFlown": "Undisclosed Gender Flown",
    "totalPilots": "Registered Pilots",
    "totalFlights": "Flights Provided",
}


class ApiError(RuntimeError):
    pass


def prompt_credentials() -> tuple[str, str]:
    email = input("Email: ").strip()
    if not email:
        raise SystemExit("Email is required.")
    password = getpass("Password: ")
    if not password:
        raise SystemExit("Password is required.")
    return email, password


def prompt_date(label: str) -> date:
    raw = input(f"{label} (YYYY-MM-DD): ").strip()
    if not raw:
        raise SystemExit(f"{label} is required.")
    try:
        return date_parser.parse(raw).date()
    except (ValueError, OverflowError, TypeError) as exc:
        raise SystemExit(f"Could not parse {label}: {raw}") from exc


def sleep_politely() -> None:
    time.sleep(REQUEST_PAUSE_SEC)


def parse_event_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return date_parser.parse(text).date()
    except (ValueError, OverflowError, TypeError):
        return None


def chapter_label(event: dict[str, Any]) -> str:
    division = nested_get(event, "chapterDivision") or {}
    name = nested_get(division, "name") if isinstance(division, dict) else ""
    number = nested_get(event, "chapterDivisionNumber")
    if number is None and isinstance(division, dict):
        number = nested_get(division, "number") or ""
    return f"{name or ''} {number or ''}".strip()


def event_type_label(event: dict[str, Any]) -> str:
    event_type = nested_get(event, "eventType") or {}
    if isinstance(event_type, dict):
        return str(nested_get(event_type, "description") or nested_get(event, "eventTypeId") or "")
    return str(event_type or nested_get(event, "eventTypeId") or "")


def nested_get(row: dict[str, Any], *keys: str) -> Any:
    if not isinstance(row, dict):
        return None
    lookup = {str(name).lower(): value for name, value in row.items()}
    for key in keys:
        if key.lower() in lookup and lookup[key.lower()] not in (None, ""):
            return lookup[key.lower()]
    for key in keys:
        if key.lower() in lookup:
            return lookup[key.lower()]
    return None


def unwrap_envelope(payload: Any) -> Any:
    """The EAA API often wraps payloads as { header, data }."""
    if not isinstance(payload, dict):
        return payload
    lookup = {str(key).lower(): value for key, value in payload.items()}
    if "header" in lookup and "data" in lookup:
        return lookup["data"]
    if "header" in lookup and "result" in lookup:
        return lookup["result"]
    return payload


def payload_shape(payload: Any) -> str:
    if isinstance(payload, dict):
        return "{" + ", ".join(f"{key}:{type(value).__name__}" for key, value in payload.items()) + "}"
    if isinstance(payload, list):
        return f"list[{len(payload)}]"
    return type(payload).__name__


def extract_rows_and_total(payload: Any, context: str) -> tuple[list[Any], int | None]:
    if payload is None:
        return [], 0
    if isinstance(payload, list):
        return payload, len(payload)
    if not isinstance(payload, dict):
        raise ApiError(f"Unexpected {context} type {type(payload).__name__}.")

    lookup = {str(key).lower(): value for key, value in payload.items()}
    inner = lookup.get("data", lookup.get("result", lookup.get("value")))
    if isinstance(inner, list):
        payload = {"list": inner, "paginationParams": lookup.get("paginationparams") or {}}
        lookup = {str(key).lower(): value for key, value in payload.items()}
    elif isinstance(inner, dict) and not any(
        str(key).lower() in {"list", "items", "events"} for key in payload
    ):
        payload = inner
        lookup = {str(key).lower(): value for key, value in payload.items()}

    chunk = lookup.get("list", lookup.get("items", lookup.get("events")))
    pagination = lookup.get("paginationparams") or lookup.get("pagination") or {}
    total = None
    if isinstance(pagination, dict):
        page_lookup = {str(key).lower(): value for key, value in pagination.items()}
        total = page_lookup.get("totalitems", page_lookup.get("totalcount", page_lookup.get("total")))
    if chunk is None:
        raise ApiError(f"Unexpected {context} response. Shape: {payload_shape(payload)}")
    if not isinstance(chunk, list):
        raise ApiError(
            f"Unexpected {context} list type {type(chunk).__name__}. Shape: {payload_shape(payload)}"
        )
    return chunk, None if total is None else int(total)


class EaaClient:
    def __init__(self, access_token: str) -> None:
        self.session = requests.Session()
        self.session.headers.update(
            {
                "Authorization": f"Bearer {access_token}",
                "Accept": "application/json",
                "Content-Type": "application/json",
                "Origin": SITE_ORIGIN,
                "Referer": f"{SITE_ORIGIN}/",
            }
        )

    def _handle(self, response: requests.Response) -> Any:
        if response.status_code == 401:
            raise ApiError("Session expired or login was rejected (HTTP 401).")
        if not response.ok:
            snippet = response.text[:400].replace("\n", " ")
            raise ApiError(f"HTTP {response.status_code} for {response.url}: {snippet}")
        if not response.content:
            return None
        try:
            return unwrap_envelope(response.json())
        except json.JSONDecodeError as exc:
            raise ApiError(f"Non-JSON response from {response.url}") from exc

    def get(self, path: str) -> Any:
        sleep_politely()
        return self._handle(self.session.get(API_BASE + path, timeout=60))

    def post(self, path: str, body: dict[str, Any]) -> Any:
        sleep_politely()
        return self._handle(self.session.post(API_BASE + path, json=body, timeout=60))

    def paginated_post(self, path: str, filters: dict[str, Any], page_size: int) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        page = 1
        total = None
        while True:
            payload = self.post(
                path,
                {
                    "filters": filters,
                    "pageSize": page_size,
                    "currentPage": page,
                },
            )
            if not payload:
                break
            chunk, page_total = extract_rows_and_total(payload, path)
            rows.extend(chunk)
            if page_total is not None:
                total = page_total
            if not chunk:
                break
            if total is not None and len(rows) >= int(total):
                break
            if len(chunk) < page_size:
                break
            page += 1
        return rows

    def get_optional(self, path: str) -> Any:
        try:
            return self.get(path)
        except ApiError as exc:
            print(f"    skip {path}: {exc}")
            return None


def recaptcha_token(page: Any) -> str:
    page.wait_for_function(
        "() => window.grecaptcha && window.grecaptcha.execute",
        timeout=30000,
    )
    token = page.evaluate(
        """async (siteKey) => {
            return await window.grecaptcha.execute(siteKey, { action: "LoginSubmit" });
        }""",
        RECAPTCHA_SITE_KEY,
    )
    if not token:
        raise ApiError("Could not obtain a reCAPTCHA token from the login page.")
    return str(token)


def login_via_api_in_browser(page: Any, email: str, password: str) -> str | None:
    token = recaptcha_token(page)
    response = requests.post(
        API_BASE + "Auth/login",
        headers={
            "Accept": "application/json",
            "Content-Type": "application/json",
            "Origin": SITE_ORIGIN,
            "Referer": LOGIN_URL,
        },
        json={
            "userName": email,
            "password": password,
            "recaptchaToken": token,
        },
        timeout=60,
    )
    if not response.ok:
        return None
    body = response.json()
    access = body.get("accessToken")
    return str(access) if access else None


def login_via_form(page: Any, email: str, password: str) -> str | None:
    user_box = page.locator("#userName")
    user_box.wait_for(state="visible", timeout=30000)
    user_box.fill(email)
    page.locator("#password").fill(password)
    page.locator('button[type="submit"]').click()
    try:
        page.wait_for_function(
            "() => Boolean(localStorage.getItem('accessToken'))",
            timeout=45000,
        )
    except PlaywrightTimeout:
        return None
    return page.evaluate("() => localStorage.getItem('accessToken')")


def login(email: str, password: str) -> str:
    print("Opening the EAA login page (required for reCAPTCHA)...")
    last_error = "Login failed."
    with sync_playwright() as playwright:
        for headed in (False, True):
            browser = playwright.chromium.launch(headless=not headed)
            context = browser.new_context()
            page = context.new_page()
            try:
                page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=60000)
                if headed:
                    print("A browser window opened. Completing login there...")
                token = None
                try:
                    token = login_via_api_in_browser(page, email, password)
                except Exception as exc:  # noqa: BLE001
                    last_error = str(exc)
                if not token:
                    token = login_via_form(page, email, password)
                if token:
                    print("Logged in.")
                    return token
                last_error = page.locator("body").inner_text()[:300]
            except Exception as exc:  # noqa: BLE001 — surface login problems clearly
                last_error = str(exc)
            finally:
                browser.close()
    raise ApiError(
        "Could not log in. Check the email and password, then try again. "
        f"Last error: {last_error}"
    )


def list_past_events(client: EaaClient) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    page = 1
    total = None
    while True:
        print(f"Fetching past events page {page}...")
        payload = client.post(
            "Event/getPastAndCurrentEvents",
            {
                "currentPage": page,
                "pageSize": EVENT_PAGE_SIZE,
                "filters": {
                    "hostChapter": "",
                    "eventTitle": "",
                    "eventType": "",
                },
            },
        )
        chunk, page_total = extract_rows_and_total(payload, "past-events")
        rows.extend(chunk)
        if page_total is not None:
            total = page_total
        if not chunk:
            break
        if total is not None and len(rows) >= int(total):
            break
        if len(chunk) < EVENT_PAGE_SIZE:
            break
        page += 1
    return rows


def in_date_range(event: dict[str, Any], start: date, end: date) -> bool:
    event_date = parse_event_date(
        nested_get(event, "eventStartDate") or nested_get(event, "eventEndDate")
    )
    if event_date is None:
        return False
    return start <= event_date <= end


def is_cancelled(event: dict[str, Any]) -> bool:
    status = nested_get(event, "eventStatusId")
    if status == STATUS_CANCELLED:
        return True
    return str(nested_get(event, "eventStatus") or "").upper() == "CANCELLED"


def is_submitted(event: dict[str, Any]) -> bool:
    status = nested_get(event, "eventStatusId")
    if status == STATUS_COMPLETED:
        return True
    return str(nested_get(event, "eventStatus") or "").upper() == "COMPLETED"


def event_meta(event: dict[str, Any]) -> dict[str, Any]:
    event_date = parse_event_date(nested_get(event, "eventStartDate"))
    status_id = nested_get(event, "eventStatusId")
    return {
        "event_date": event_date.isoformat() if event_date else "",
        "event_title": nested_get(event, "eventTitle") or "",
        "event_id": nested_get(event, "eventId"),
        "event_link_id": nested_get(event, "eventLinkId"),
        "chapter": chapter_label(event),
        "event_type": event_type_label(event),
        "event_status": STATUS_NAMES.get(status_id, status_id),
    }


def fetch_final_report(client: EaaClient, event: dict[str, Any], index: int, total: int) -> dict[str, Any]:
    event_id = nested_get(event, "eventId")
    title = nested_get(event, "eventTitle") or event_id
    print(f"[{index}/{total}] Final report: {title}")
    hq = client.get(f"FinalReport/GetSubmissionToHeadquarters/{event_id}")
    age = client.get(f"FinalReport/getAgeReport/{event_id}")
    pilots = client.get(f"FinalReport/GetPilotFlightsReport/{event_id}")
    if isinstance(age, dict):
        age, _ = extract_rows_and_total(age, "age report")
    if isinstance(pilots, dict):
        pilots, _ = extract_rows_and_total(pilots, "pilot flights")
    if isinstance(hq, dict) and nested_get(hq, "totalYoungs", "totalYoungsFlown") is None:
        inner = nested_get(hq, "data", "result")
        if isinstance(inner, dict):
            hq = inner
    parents: dict[int, list[dict[str, Any]]] = {}
    for option, sheet_name, _label in PARENT_FILTERS:
        print(f"    parents ({sheet_name})...")
        parents[option] = client.paginated_post(
            "FinalReport/GetRegisteredParentsReport",
            {"eventId": event_id, "filterOption": option},
            REPORT_PAGE_SIZE,
        )
    print("    volunteers...")
    volunteers = client.paginated_post(
        "FinalReport/GetVolunteersReport",
        {"eventId": event_id},
        REPORT_PAGE_SIZE,
    )
    print("    youth ages...")
    event_date = parse_event_date(nested_get(event, "eventStartDate"))
    parent_ids = collect_parent_ids(parents)
    age_lookup = age_lookup_from_youth_rows(
        fetch_youth_rows_with_age(client, parent_ids),
        event_date,
    )
    attach_ages_to_parents(parents, age_lookup, event_date)
    return {
        "event": event,
        "hq": hq if isinstance(hq, dict) else {},
        "age": age if isinstance(age, list) else [],
        "pilots": pilots if isinstance(pilots, list) else [],
        "parents": parents,
        "volunteers": volunteers,
    }


def flatten_pilot_ages(pilot: dict[str, Any]) -> dict[str, Any]:
    by_age = {item.get("age"): item.get("totalYoungs") for item in (pilot.get("youngsByAge") or [])}
    columns = {f"age_{age}": by_age.get(age, 0) for age in range(8, 18)}
    columns["ages"] = ", ".join(
        f"{age}:{by_age.get(age, 0)}" for age in range(8, 18) if by_age.get(age)
    )
    return columns


def normalize_name_part(value: Any) -> str:
    text = str(value or "").casefold().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def format_age_value(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def row_youth_names(row: dict[str, Any]) -> tuple[Any, Any]:
    first = nested_get(row, "kidFirstName", "youthFirstName", "youngFirstName", "firstName")
    last = nested_get(row, "kidLastName", "youthLastName", "youngLastName", "lastName")
    young = nested_get(row, "young")
    if isinstance(young, dict):
        first = first or nested_get(young, "firstName")
        last = last or nested_get(young, "lastName")
    full = nested_get(row, "youngFullName", "youthFullName", "kidFullName")
    if (not first or not last) and isinstance(full, str):
        parts = full.strip().split()
        if parts:
            first = first or parts[0]
            last = last or (parts[-1] if len(parts) > 1 else last)
    return first, last


def youth_age_from_row(row: dict[str, Any], on_date: date | None = None) -> str:
    young = nested_get(row, "young") if isinstance(nested_get(row, "young"), dict) else {}
    age = nested_get(row, "age", "youngAge", "kidAge")
    if age in (None, "") and young:
        age = nested_get(young, "age")
    formatted = format_age_value(age)
    if formatted:
        return formatted
    birthday = nested_get(row, "birthday", "birthDate", "dateOfBirth")
    if birthday in (None, "") and young:
        birthday = nested_get(young, "birthday", "birthDate", "dateOfBirth")
    if birthday in (None, "") or on_date is None:
        return ""
    born = parse_event_date(birthday)
    if born is None:
        return ""
    years = on_date.year - born.year - ((on_date.month, on_date.day) < (born.month, born.day))
    if years < 0 or years > 120:
        return ""
    return str(years)


def youth_key(row: dict[str, Any]) -> tuple[str, str]:
    first, last = row_youth_names(row)
    return (normalize_name_part(first), normalize_name_part(last))


def youth_display(row: dict[str, Any]) -> dict[str, str]:
    first, last = row_youth_names(row)
    first = str(first or "").strip()
    last = str(last or "").strip()
    parent_first = str(nested_get(row, "parentFirstName") or "").strip()
    parent_last = str(nested_get(row, "parentLastName") or "").strip()
    email = str(nested_get(row, "parentAddress", "parentEmail", "email") or "").strip()
    mailing = str(nested_get(row, "parentMailingAddress", "mailingAddress") or "").strip()
    return {
        "youth_first_name": first,
        "youth_last_name": last,
        "youth_name": f"{first} {last}".strip(),
        "age": youth_age_from_row(row),
        "parent_name": f"{parent_first} {parent_last}".strip(),
        "parent_email": email,
        "parent_mailing_address": mailing,
    }


def unique_join(values: list[str]) -> str:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        text = str(value or "").strip()
        marker = re.sub(r"\s+", " ", text.casefold())
        if not text or marker in seen:
            continue
        seen.add(marker)
        ordered.append(text)
    return "; ".join(ordered)


def payload_as_youth_rows(payload: Any, context: str) -> list[dict[str, Any]]:
    if not payload:
        return []
    if isinstance(payload, list):
        return [row for row in payload if isinstance(row, dict)]
    if not isinstance(payload, dict):
        return []
    try:
        rows, _ = extract_rows_and_total(payload, context)
        return [row for row in rows if isinstance(row, dict)]
    except ApiError:
        pass
    for key in (
        "youngs",
        "youths",
        "youngCheckIn",
        "checkInYouths",
        "youngByPilotInEvent",
        "listYoungCheckIn",
        "young",
    ):
        inner = nested_get(payload, key)
        if isinstance(inner, list):
            return [row for row in inner if isinstance(row, dict)]
    return []


def collect_parent_ids(parents: dict[int, list[dict[str, Any]]]) -> list[Any]:
    parent_ids: list[Any] = []
    seen_parent: set[str] = set()
    for rows in parents.values():
        for row in rows:
            parent_id = nested_get(row, "parentId")
            if parent_id not in (None, "") and str(parent_id) not in seen_parent:
                seen_parent.add(str(parent_id))
                parent_ids.append(parent_id)
    return parent_ids


def fetch_youth_rows_with_age(client: EaaClient, parent_ids: list[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for parent_id in parent_ids:
        payload = client.get_optional(
            f"EventRegistrationConfirmation/GetRegisterDetails/{parent_id}"
        )
        rows.extend(payload_as_youth_rows(payload, "register-details"))
    return rows


def age_lookup_from_youth_rows(
    rows: list[dict[str, Any]], on_date: date | None = None
) -> dict[str, Any]:
    collected: dict[tuple[str, str], list[str]] = defaultdict(list)
    by_id: dict[str, str] = {}
    for row in rows:
        age = youth_age_from_row(row, on_date)
        young_id = nested_get(row, "youngId")
        if young_id not in (None, "") and age:
            by_id[str(young_id)] = age
        key = youth_key(row)
        if any(key) and age:
            collected[key].append(age)
    return {
        "by_name": {key: unique_join(ages) for key, ages in collected.items()},
        "by_id": by_id,
    }


def attach_ages_to_parents(
    parents: dict[int, list[dict[str, Any]]],
    lookup: dict[str, Any],
    on_date: date | None = None,
) -> None:
    by_name = lookup.get("by_name") or {}
    by_id = lookup.get("by_id") or {}
    for rows in parents.values():
        for row in rows:
            current = youth_age_from_row(row, on_date)
            if current:
                row["age"] = current
                continue
            young_id = nested_get(row, "youngId")
            if young_id not in (None, "") and str(young_id) in by_id:
                row["age"] = by_id[str(young_id)]
                continue
            key = youth_key(row)
            if key in by_name:
                row["age"] = by_name[key]


def tally_child_flights(reports: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[tuple[str, str], dict[str, Any]] = {}
    for report in reports:
        meta = event_meta(report["event"])
        event_label = f"{meta['event_date']} {meta['event_title']}".strip()
        for row in report["parents"].get(2, []):
            key = youth_key(row)
            if not any(key):
                continue
            display = youth_display(row)
            bucket = buckets.get(key)
            if bucket is None:
                bucket = {
                    "youth_first_name": display["youth_first_name"],
                    "youth_last_name": display["youth_last_name"],
                    "youth_name": display["youth_name"],
                    "parent_names": [],
                    "parent_emails": [],
                    "parent_addresses": [],
                    "ages": [],
                    "events": [],
                }
                buckets[key] = bucket
            if event_label and event_label not in bucket["events"]:
                bucket["events"].append(event_label)
            bucket["parent_names"].append(display["parent_name"])
            bucket["parent_emails"].append(display["parent_email"])
            bucket["parent_addresses"].append(display["parent_mailing_address"])
            bucket["ages"].append(display["age"])
    rows = []
    for bucket in buckets.values():
        rows.append(
            {
                "youth_name": bucket["youth_name"],
                "youth_first_name": bucket["youth_first_name"],
                "youth_last_name": bucket["youth_last_name"],
                "age": unique_join(bucket["ages"]),
                "parent_name": unique_join(bucket["parent_names"]),
                "parent_email": unique_join(bucket["parent_emails"]),
                "parent_mailing_address": unique_join(bucket["parent_addresses"]),
                "flights_in_range": len(bucket["events"]),
                "events": "; ".join(bucket["events"]),
            }
        )
    rows.sort(key=lambda item: (-item["flights_in_range"], item["youth_last_name"], item["youth_first_name"]))
    return rows


def autosize(worksheet) -> None:
    for column in worksheet.columns:
        letter = get_column_letter(column[0].column)
        width = 12
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            width = min(max(width, len(value) + 2), 50)
        worksheet.column_dimensions[letter].width = width


def write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    sheet = workbook.create_sheet(title[:31])
    header_font = Font(bold=True)
    for col_index, name in enumerate(columns, start=1):
        cell = sheet.cell(1, col_index, name)
        cell.font = header_font
    for row_index, row in enumerate(rows, start=2):
        for col_index, name in enumerate(columns, start=1):
            sheet.cell(row_index, col_index, row.get(name, ""))
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        sheet.freeze_panes = "A2"
    autosize(sheet)


def bool_label(value: Any) -> str:
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return "" if value is None else str(value)


def build_workbook(reports: list[dict[str, Any]], start: date, end: date) -> Workbook:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    event_rows = []
    hq_rows = []
    age_rows = []
    pilot_rows = []
    volunteer_rows = []
    parent_rows: dict[int, list[dict[str, Any]]] = defaultdict(list)

    for report in reports:
        event = report["event"]
        meta = event_meta(event)
        event_rows.append(
            {
                **meta,
                "number_registered": nested_get(event, "numberYoungRegistered"),
                "number_flown": nested_get(event, "numberYoungFlown"),
                "event_submitted": "Yes" if is_submitted(event) else "No",
            }
        )
        hq_rows.append(
            {**meta, **{HQ_LABELS[key]: nested_get(report["hq"] or {}, key) for key in HQ_LABELS}}
        )
        for age in report["age"]:
            age_rows.append(
                {
                    **meta,
                    "age": age.get("age"),
                    "total_female": age.get("totalGirls", age.get("totalFemale")),
                    "total_male": age.get("totalBoys", age.get("totalMale")),
                    "total_undisclosed": age.get("totalNS", age.get("totalNoSpecified")),
                }
            )
        for pilot in report["pilots"]:
            pilot_rows.append(
                {
                    **meta,
                    "pilot_name": pilot.get("fullName"),
                    "flights": pilot.get("flights"),
                    "youth_flown": pilot.get("totalYoungs"),
                    "girls": pilot.get("totalGirls"),
                    "boys": pilot.get("totalBoys"),
                    "undisclosed_gender": pilot.get("totalNS"),
                    "first_flight": pilot.get("firstFlight"),
                    "return_ye": pilot.get("noFirstFlight"),
                    **flatten_pilot_ages(pilot),
                }
            )
        for option, _sheet, _label in PARENT_FILTERS:
            for row in report["parents"].get(option, []):
                parent_rows[option].append(
                    {
                        **meta,
                        "parent_first_name": row.get("parentFirstName"),
                        "parent_last_name": row.get("parentLastName"),
                        "youth_first_name": row.get("kidFirstName"),
                        "youth_last_name": nested_get(row, "kidLastName"),
                        "age": youth_age_from_row(row),
                        "parent_email": nested_get(row, "parentAddress"),
                        "parent_mailing_address": row.get("parentMailingAddress"),
                        "waiver_signed": bool_label(
                            nested_get(row, "waiverSignatureFromMobile", "waiverSigned")
                        ),
                    }
                )
        for volunteer in report["volunteers"]:
            volunteer_rows.append(
                {
                    **meta,
                    "type": volunteer.get("type"),
                    "first_name": volunteer.get("firstName"),
                    "last_name": volunteer.get("lastName"),
                    "email": volunteer.get("email"),
                }
            )

    event_columns = [
        "event_date",
        "chapter",
        "event_title",
        "event_type",
        "number_registered",
        "number_flown",
        "event_submitted",
        "event_status",
        "event_id",
        "event_link_id",
    ]
    shared_event = ["event_date", "event_title", "event_id", "event_link_id", "chapter"]
    write_sheet(workbook, "Events", event_rows, event_columns)
    write_sheet(
        workbook,
        "Headquarters",
        hq_rows,
        shared_event + list(HQ_LABELS.values()),
    )
    write_sheet(
        workbook,
        "Age",
        age_rows,
        shared_event + ["age", "total_female", "total_male", "total_undisclosed"],
    )
    write_sheet(
        workbook,
        "Pilot Flights",
        pilot_rows,
        shared_event
        + [
            "pilot_name",
            "flights",
            "youth_flown",
            "girls",
            "boys",
            "undisclosed_gender",
            "first_flight",
            "return_ye",
            "ages",
        ]
        + [f"age_{age}" for age in range(8, 18)],
    )
    parent_columns = shared_event + [
        "parent_first_name",
        "parent_last_name",
        "youth_first_name",
        "youth_last_name",
        "age",
        "parent_email",
        "parent_mailing_address",
        "waiver_signed",
    ]
    for option, sheet_name, _label in PARENT_FILTERS:
        write_sheet(workbook, sheet_name, parent_rows[option], parent_columns)
    write_sheet(
        workbook,
        "Volunteers",
        volunteer_rows,
        shared_event + ["type", "first_name", "last_name", "email"],
    )
    child_rows = tally_child_flights(reports)
    write_sheet(
        workbook,
        "Child Flights",
        child_rows,
        [
            "youth_name",
            "youth_first_name",
            "youth_last_name",
            "age",
            "parent_name",
            "parent_email",
            "parent_mailing_address",
            "flights_in_range",
            "events",
        ],
    )
    workbook.properties.title = f"EAA final reports {start.isoformat()} to {end.isoformat()}"
    return workbook


def main() -> int:
    print("EAA Final Report Excel export")
    print("Credentials are used only for this run and are not saved.\n")
    email, password = prompt_credentials()
    start = prompt_date("Start date")
    end = prompt_date("End date")
    if end < start:
        print("End date must be on or after the start date.")
        return 1

    try:
        token = login(email, password)
        client = EaaClient(token)
        events = list_past_events(client)
        selected = [
            event
            for event in events
            if in_date_range(event, start, end)
            and not is_cancelled(event)
            and is_submitted(event)
        ]
        print(
            f"Found {len(events)} past/current events; "
            f"{len(selected)} in range {start.isoformat()} to {end.isoformat()} "
            f"(cancelled and unsubmitted skipped)."
        )
        if not selected:
            print("No matching events. No workbook was written.")
            return 0

        reports = [
            fetch_final_report(client, event, index, len(selected))
            for index, event in enumerate(selected, start=1)
        ]
        workbook = build_workbook(reports, start, end)
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        filename = REPORTS_DIR / f"eaa_final_reports_{start.isoformat()}_to_{end.isoformat()}.xlsx"
        workbook.save(filename)
        print(f"Wrote {filename}")
        return 0
    except ApiError as exc:
        print(f"Error: {exc}")
        return 1
    except KeyboardInterrupt:
        print("\nCancelled.")
        return 130


if __name__ == "__main__":
    sys.exit(main())
